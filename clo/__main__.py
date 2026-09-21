"""Command-line runner. All portfolio processing and workbook rendering are local."""
from __future__ import annotations
import argparse
import copy
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime, timezone

from .inputs import FIELDS, ValidationError, load_portfolio, read_json, summarize, validate_config, digest, sha256
from .methodology import supplemental_tests, SOURCE_META, credit_quality_metrics, SP_WARF_FACTORS, SP_WARF_SOURCE
from .fitch import fitch_metrics

ROOT=Path(__file__).resolve().parents[1]

# These are exactly the scalar inputs exposed by the workbook builder. Structural
# JSON objects, mappings and vocabularies are edited in the configuration file.
SCALAR_ASSUMPTIONS = {
    'as_of_date','legal_final_date','currency','benchmark','base_senior_pct','equity_pct',
    'senior_coupon_bps','junior_coupon_bps','senior_floor_pct','junior_floor_pct',
    'senior_fee_bps','fixed_fee_per_period','subordinated_fee_bps','oc_trigger','ic_trigger',
    'recovery_lag_quarters','annual_prepayment_rate','reinvestment_end_date',
    'reinvestment_price','reinvestment_spread_bps','reinvestment_recovery',
    'reinvestment_maturity_years','reinvestment_floor_bps','principal_support_interest',
    'final_liquidation_price','oc_recovery_credit','reinvestment_stop_on_trigger',
    'base_default_rate','base_scenario_id','agency_sdr_aaa','agency_sdr_source',
}
SCENARIO_ASSUMPTIONS = {'recovery_multiplier','annual_prepayment_rate'}

def write_json(path, obj):
    Path(path).write_text(json.dumps(obj,indent=2,allow_nan=False)+'\n',encoding='utf-8')

def diagnostic_value(value):
    """Keep invalid values legible without emitting nonstandard JSON numbers."""
    if isinstance(value,float) and not math.isfinite(value): return repr(value)
    if isinstance(value,dict): return {str(k):diagnostic_value(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)): return [diagnostic_value(v) for v in value]
    if isinstance(value,(date,datetime)): return value.isoformat()
    return value

def write_diagnostic(path, obj):
    """A diagnostic I/O failure must not mask the original failure or touch results."""
    try:
        write_json(path,diagnostic_value(obj))
        return True
    except (OSError,TypeError,ValueError) as exc:
        print(f'Could not save diagnostic report: {exc}',file=sys.stderr)
        return False

def paths_overlap(first, second):
    a,b=Path(first).resolve(),Path(second).resolve()
    if os.path.normcase(str(a))==os.path.normcase(str(b)): return True
    try: return a.exists() and b.exists() and a.samefile(b)
    except OSError: return False

def validate_output_paths(args):
    """Outputs, including sidecars, may never overwrite an input or each other."""
    destinations=[args.output,args.output.with_suffix('.results.json'),args.output.with_suffix('.validation.json')]
    sources=[p for p in (args.portfolio,args.config,args.assumptions_from,getattr(args,'fitch_inputs',None)) if p is not None]
    for position,target in enumerate(destinations):
        for source in sources:
            if paths_overlap(target,source):
                raise ValidationError([{'field':'output','message':f'Output path would overwrite input: {target}. Choose a different output name.'}])
        if any(paths_overlap(target,other) for other in destinations[:position]):
            raise ValidationError([{'field':'output','message':'Workbook and sidecar output paths must be distinct.'}])

def publish(staged_outputs, temp_dir):
    """Stage all artifacts, back up existing destinations, and roll back failures."""
    backups={}
    for i,(_,destination) in enumerate(staged_outputs):
        if destination.exists():
            backup=Path(temp_dir)/f'previous_{i}.bin'
            shutil.copy2(destination,backup)
            backups[destination]=backup
    replaced=[]
    try:
        for staged,destination in staged_outputs:
            os.replace(staged,destination)
            replaced.append(destination)
    except Exception:
        restoration_errors=[]
        for destination in reversed(replaced):
            try:
                if destination in backups:
                    # Copying the retained backup also works when a mocked or
                    # platform-specific rename failure affects the destination.
                    shutil.copy2(backups[destination],destination)
                else:
                    destination.unlink(missing_ok=True)
            except OSError as exc:
                restoration_errors.append(f'{destination}: {exc}')
        if restoration_errors:
            raise RuntimeError('Publication failed and restoration requires attention: '+'; '.join(restoration_errors))
        raise

def finalize_xlsx(path):
    """Add Excel recalculation flags missing from the artifact export API.

    Only workbook.xml's calcPr is altered. Cells, formulas, caches, relationships
    and all other package members remain byte-for-byte unchanged.
    """
    path=Path(path)
    with zipfile.ZipFile(path,'r') as source:
        workbook_xml=source.read('xl/workbook.xml').decode('utf-8')
        ET.fromstring(workbook_xml)  # Reject malformed XML before changing the package.
        opening=re.search(r'<((?:[A-Za-z_][\w.-]*:)?workbook)\b',workbook_xml)
        if opening is None: raise ValueError('Excel package has no workbook XML root')
        prefix=opening.group(1).removesuffix('workbook')
        calc_tag=prefix+'calcPr'
        tag_pattern=re.compile(r'<'+re.escape(calc_tag)+r'\b[^>]*(?:/\s*>|>\s*</'+re.escape(calc_tag)+r'\s*>)')
        existing=tag_pattern.search(workbook_xml)
        if existing:
            replacement=existing.group(0)
            for key,value in [('calcMode','auto'),('fullCalcOnLoad','1'),('forceFullCalc','1')]:
                attr=re.compile(r'\s'+key+r'\s*=\s*([\"\']).*?\1')
                if attr.search(replacement): replacement=attr.sub(f' {key}="{value}"',replacement)
                else: replacement=replacement.replace('<'+calc_tag,'<'+calc_tag+f' {key}="{value}"',1)
            workbook_xml=workbook_xml[:existing.start()]+replacement+workbook_xml[existing.end():]
        else:
            replacement=f'<{calc_tag} calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
            after_calc=('oleSize','customWorkbookViews','pivotCaches','smartTagPr','smartTagTypes',
                        'webPublishing','fileRecoveryPr','webPublishObjects','extLst')
            following=re.search(r'<'+re.escape(prefix)+r'(?:'+'|'.join(after_calc)+r')\b',workbook_xml)
            at=following.start() if following else workbook_xml.rfind('</'+opening.group(1)+'>')
            if at<0: raise ValueError('Excel package has no closing workbook XML tag')
            workbook_xml=workbook_xml[:at]+replacement+workbook_xml[at:]
        ET.fromstring(workbook_xml)
        revised=path.with_name(path.name+'.recalc.tmp')
        with zipfile.ZipFile(revised,'w') as target:
            for entry in source.infolist():
                target.writestr(entry,workbook_xml.encode('utf-8') if entry.filename=='xl/workbook.xml' else source.read(entry.filename))
    os.replace(revised,path)

def validate_external_sdr(config,portfolio):
    if config.get('agency_sdr_aaa') is None: return
    expected={
        'agency_sdr_portfolio_hash':portfolio['source_hash'],
        'agency_sdr_as_of':config['as_of_date'],
        'agency_sdr_criteria_revision':SOURCE_META['republication'],
        'agency_sdr_credit_hash':digest({'facilities':portfolio['facilities'],
                                      'as_of_date':config['as_of_date'],
                                      'criteria_revision':SOURCE_META['republication']}),
    }
    issues=[]
    for key,value in expected.items():
        actual=config.get(key)
        if key.endswith('_hash') and isinstance(actual,str): actual=actual.lower()
        if actual!=value:
            issues.append({'field':key,'message':'External SDR requires matching portfolio, normalized credit inputs, as-of date and implemented criteria revision. Update the approved external SDR or clear agency_sdr_aaa.'})
    for key in ('agency_sdr_source','agency_sdr_model_version'):
        if not isinstance(config.get(key),str) or not config[key].strip():
            issues.append({'field':key,'message':'External SDR requires nonblank source and model-version provenance.'})
    if issues: raise ValidationError(issues)

def import_assumptions(config,path):
    """Import designated literal cells; ignore display dates and run snapshots."""
    import openpyxl
    workbook=openpyxl.load_workbook(path,data_only=False,read_only=True)
    updated=copy.deepcopy(config)
    def invalid(field,message):
        raise ValidationError([{'field':field,'message':message}])
    def literal(value,field):
        if isinstance(value,str) and value.lstrip().startswith('='):
            invalid(field,'Editable inputs must be literal values; replace formula with a value')
        if isinstance(value,datetime): return value.date().isoformat()
        if isinstance(value,date): return value.isoformat()
        return value
    def finite_number(value):
        return not isinstance(value,bool) and isinstance(value,(int,float)) and math.isfinite(value)
    def unchanged(value,original):
        if isinstance(value,bool) or isinstance(original,bool):
            return type(value) is type(original) and value==original
        return value==original
    try:
        if not {'Assumptions','Stress Inputs'}<=set(workbook.sheetnames):
            invalid('workbook','Not a model workbook: missing editable input tabs')
        required_keys = SCALAR_ASSUMPTIONS | {
            f"scenario.{scenario['id']}.{field}"
            for scenario in config['scenarios'] for field in SCENARIO_ASSUMPTIONS
        }
        seen=set()
        for row in workbook['Assumptions'].iter_rows(min_row=8,values_only=True):
            key=row[2] if len(row)>2 else None
            value=row[3] if len(row)>3 else None
            if key is None or key=='':
                if value is not None: invalid('assumption_key','An edited value has no assumption key')
                continue
            if not isinstance(key,str): invalid('assumption_key','Assumption keys must be text')
            if key in seen: invalid(key,'Duplicate assumption key')
            seen.add(key)
            value=literal(value,key)
            if key.startswith('scenario.'):
                pieces=key[len('scenario.'):].rsplit('.',1)
                if len(pieces)!=2 or pieces[1] not in SCENARIO_ASSUMPTIONS:
                    invalid(key,'Unknown editable scenario assumption')
                sid,field=pieces
                matches=[s for s in updated['scenarios'] if s['id']==sid]
                if not matches: invalid(key,'Scenario IDs differ from configuration')
                if not unchanged(value,matches[0].get(field)): matches[0][field]=value
            elif key in SCALAR_ASSUMPTIONS and key in updated:
                # Preserve unchanged numeric representations (0.0 vs 0) so an
                # untouched workbook imports to the identical configuration hash.
                if not unchanged(value,updated[key]): updated[key]=value
            else: invalid(key,'Unknown assumption key; use the matching config file')
        missing_keys = required_keys - seen
        if missing_keys:
            invalid('assumptions', 'Missing assumption rows: '+', '.join(sorted(missing_keys))+
                    '. Restore the complete table or regenerate the workbook; configuration values are not substituted.')
        curves={s['id']:{'weights':{},'rates':{}} for s in updated['scenarios']}
        for row in workbook['Stress Inputs'].iter_rows(min_row=8,values_only=True):
            sid=row[0] if row else None
            if sid is None or sid=='':
                if any(v is not None for v in row[1:5]): invalid('scenario','Stress values have no scenario ID')
                continue
            if not isinstance(sid,str) or sid not in curves: invalid('scenario',f'Unknown stress scenario {sid}')
            if len(row)<5: invalid('stress_inputs','Stress row is missing period, weight or rate cells')
            period=row[1]
            if not finite_number(period) or period!=int(period) or period<1:
                invalid('period','Stress period must be a finite positive integer')
            period=int(period)
            if period in curves[sid]['weights']: invalid('period',f'Duplicate period {sid}/{period}')
            weight,rate=literal(row[3],'default_weight'),literal(row[4],'index_rate')
            if not finite_number(weight) or weight<0: invalid('default_weight','Default weights must be finite nonnegative numbers')
            if not finite_number(rate) or not -.1<=rate<=1: invalid('index_rate','Index rates must be finite numeric fractions in -0.1..1')
            curves[sid]['weights'][period]=weight
            curves[sid]['rates'][period]=rate
        for original,scenario in zip(config['scenarios'],updated['scenarios']):
            curve=curves[scenario['id']]
            old_weights=original['default_weights']
            old_rates=original.get('index_rate_path',config['index_rate_path'])
            expected=set(range(1,max(len(old_weights),len(old_rates))+1))
            if set(curve['weights'])!=expected:
                invalid('stress_inputs',f'Missing or added stress periods: {scenario["id"]}')
            weights=[curve['weights'][i] for i in sorted(expected)]
            rates=[curve['rates'][i] for i in sorted(expected)]
            if all(v==0 for v in weights[len(old_weights):]): weights=weights[:len(old_weights)]
            if all(v==old_rates[-1] for v in rates[len(old_rates):]): rates=rates[:len(old_rates)]
            scenario['default_weights']=[old_weights[i] if i<len(old_weights) and v==old_weights[i] else v for i,v in enumerate(weights)]
            rates=[old_rates[i] if i<len(old_rates) and v==old_rates[i] else v for i,v in enumerate(rates)]
            if 'index_rate_path' in original or rates!=old_rates: scenario['index_rate_path']=rates
        return updated
    finally:
        workbook.close()

def renderer_paths():
    """Locate local bundled renderer without downloading any dependency or data."""
    deps=Path.home()/'.cache'/'codex-runtimes'/'codex-primary-runtime'/'dependencies'
    node=os.environ.get('CLO_NODE') or str(deps/'node'/'bin'/'node.exe')
    if not Path(node).is_file(): node=shutil.which('node')
    modules=Path(os.environ.get('CLO_NODE_MODULES',str(deps/'node'/'node_modules')))
    if not node or not (modules/'@oai'/'artifact-tool').exists():
        raise RuntimeError('Local Excel renderer unavailable. Set CLO_NODE and CLO_NODE_MODULES to an approved runtime containing @oai/artifact-tool. The Python engine can still run with --engine-only.')
    # The renderer resolves the approved local dependency with createRequire;
    # it neither downloads packages nor creates dependency links.
    return node,modules

def render(packet_path,output_path,qa_dir=None):
    node,modules=renderer_paths()
    env=os.environ.copy()
    env['CLO_ARTIFACT_MODULES']=str(modules)
    args=[node,str(ROOT/'build_workbook.mjs'),str(packet_path),str(output_path)]
    if qa_dir: args.extend(['--qa',str(qa_dir)])
    subprocess.run(args,check=True,env=env,cwd=ROOT)

def main(argv=None):
    parser=argparse.ArgumentParser(description='Validate and process a local CLO portfolio; produces no agency rating.')
    parser.add_argument('--portfolio',type=Path)
    parser.add_argument('--fitch-inputs',type=Path,help='Optional separately sourced Fitch facility inputs; no automatic S&P mapping')
    parser.add_argument('--config',type=Path,default=ROOT/'config'/'model.json')
    parser.add_argument('--output',type=Path,default=Path('CLO_Model.xlsx'))
    parser.add_argument('--assumptions-from',type=Path,help='Read designated edited Excel assumption/rate cells; CSV remains portfolio authority')
    parser.add_argument('--validate-only',action='store_true')
    parser.add_argument('--engine-only',action='store_true',help='Produce result JSON and validation without Excel renderer')
    parser.add_argument('--template',action='store_true',help='Produce blank reusable Excel review template')
    parser.add_argument('--qa-dir',type=Path,help='Optional local rendering/inspection evidence directory')
    args=parser.parse_args(argv)
    if args.template and args.validate_only:
        parser.error('--validate-only checks a portfolio; it cannot be combined with --template')
    args.output=args.output.resolve()
    report=args.output.with_suffix('.validation.json')
    # A collision cannot be reported to the colliding filename without destroying
    # the input. Reject it before creating directories or writing any output.
    try:
        validate_output_paths(args)
    except ValidationError as exc:
        print('Input/output path conflict: '+exc.issues[0]['message'],file=sys.stderr)
        return 2
    try:
        args.output.parent.mkdir(parents=True,exist_ok=True)
        config=read_json(args.config)
        validate_config(config)
        if args.assumptions_from: config=import_assumptions(config,args.assumptions_from)
        validate_config(config)
        if args.template:
            if args.fitch_inputs: parser.error('--fitch-inputs requires a portfolio; omit it for --template')
            portfolio={'facilities':[],'normalizations':[],'source_hash':'not run','source_rows':0,'source_name':'No portfolio loaded'}
            analytics={'total_par':0,'obligor_count':0,'obligors':[],'composition':{},'facility_count':0}
            concentration={}
            analysis={'base_run':{'periods':[],'summary':{},'checks':{}},'stress_results':[],'sensitivities':[],'boundaries':[]}
        else:
            if args.portfolio is None: parser.error('--portfolio is required unless --template is used')
            portfolio=load_portfolio(args.portfolio,config)
            analytics=summarize(portfolio['facilities'],config)
            validate_external_sdr(config,portfolio)
            fitch=fitch_metrics(portfolio['facilities'],args.fitch_inputs)
            if args.validate_only:
                write_json(report,{'status':'valid','source_rows':portfolio['source_rows'],'obligors':analytics['obligor_count'],'source_hash':portfolio['source_hash'],'normalizations':portfolio['normalizations'],'fitch_status':fitch['status']})
                print(f'Validated {portfolio["source_rows"]} facilities / {analytics["obligor_count"]} obligors. {report.name}')
                return 0
            print(f'Validated {portfolio["source_rows"]} facilities / {analytics["obligor_count"]} obligors; running fixed stress matrix.',flush=True)
            concentration=supplemental_tests(portfolio['facilities'],'AAA')
            from .cashflow import analyze
            analysis=analyze(portfolio['facilities'],config)
        packet={'schema_version':'1.0','template':args.template,'generated_at':datetime.now(timezone.utc).isoformat(),
                'config':config,'portfolio':portfolio,'analytics':analytics,'concentration':concentration,'analysis':analysis,
                'fitch':fitch_metrics([]) if args.template else fitch,
                'credit_quality':credit_quality_metrics(portfolio['facilities']) if portfolio['facilities'] else
                    {'spwarf':None,'default_rate_dispersion':None,'rating_factors':SP_WARF_FACTORS,'source':SP_WARF_SOURCE,'eligible_par':0,'excluded_par':0},
                'methodology':SOURCE_META,'config_hash':digest(config),'engine_version':'0.2.0',
                'run_id':digest({'portfolio':portfolio['source_hash'],'config':config,'fitch':None if args.template else fitch,'engine':'0.2.0'})[:16]}
        success_report={'status':'template' if args.template else 'valid','source_rows':portfolio['source_rows'],'obligors':analytics['obligor_count'],
                        'source_hash':portfolio['source_hash'],'config_hash':packet['config_hash'],'run_id':packet['run_id'],
                        'normalizations':portfolio['normalizations'],'methodology':'Public components + exploratory cash flows; AAA not determined'}
        # Everything is staged before publication. Existing workbook and result
        # JSON are restored together if any artifact replacement fails.
        with tempfile.TemporaryDirectory(prefix='clo_run_',dir=args.output.parent) as temp:
            temp=Path(temp)
            result=temp/'results.json'
            write_json(result,packet)
            staged=[(result,args.output.with_suffix('.results.json'))]
            if not args.engine_only:
                render(result,temp/'model.xlsx',args.qa_dir.resolve() if args.qa_dir else None)
                finalize_xlsx(temp/'model.xlsx')
                staged.append((temp/'model.xlsx',args.output))
            write_json(temp/'validation.json',success_report)
            staged.append((temp/'validation.json',report))
            publish(staged,temp)
        print(f'Created {args.output.name if not args.engine_only else args.output.with_suffix(".results.json").name}; run {packet["run_id"]}',flush=True)
        return 0
    except ValidationError as exc:
        write_diagnostic(report,{'status':'invalid','errors':exc.issues,'workbook_written':False})
        print(f'Input validation failed: {len(exc.issues)} error(s). See {report}',file=sys.stderr)
        return 2
    except Exception as exc:
        write_diagnostic(report,{'status':'failed','message':str(exc),'error_type':type(exc).__name__,'workbook_written':False})
        print(f'Run failed: {exc}. Check the diagnostic report; no successful run was published.',file=sys.stderr)
        return 1

if __name__=='__main__': raise SystemExit(main())
